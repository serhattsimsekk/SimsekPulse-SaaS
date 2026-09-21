from __future__ import annotations

import json
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from api.dependencies import get_db
from api.schemas import CellPermission, SpreadsheetCellChange, SpreadsheetCreate
from api.security import current_principal, require_roles
from database.models import SpreadsheetChange, SpreadsheetWorkbook
from database.session import SessionLocal
from services.security import decode_access_token
from api.routes.assets import _arvento_request, _arvento_settings

router = APIRouter(prefix="/api/v1/spreadsheets", tags=["Enterprise Web Excel"])


def _can_edit(protection: dict, role: str, sheet: str, cell: str) -> bool:
    for item in protection.get("cells", []):
        if item.get("sheet_name") == sheet and item.get("cell_ref") == cell:
            return role in item.get("roles", []) and bool(item.get("editable", False))
    return role in protection.get("editable_roles", ["admin"])


@router.post("/workbooks", status_code=201)
def create_workbook(payload: SpreadsheetCreate, principal: dict = Depends(current_principal), db: Session = Depends(get_db)):
    workbook = SpreadsheetWorkbook(
        tenant_id=principal["tenant_id"], owner_user_id=principal["sub"], name=payload.name,
        workbook_json=json.dumps(payload.workbook, ensure_ascii=False),
        protection_json=json.dumps(payload.protection, ensure_ascii=False),
    )
    db.add(workbook)
    db.commit()
    db.refresh(workbook)
    return {"id": workbook.id, "name": workbook.name, "version": workbook.version}


@router.get("/workbooks")
def list_workbooks(principal: dict = Depends(current_principal), db: Session = Depends(get_db)):
    return list(db.scalars(select(SpreadsheetWorkbook).where(SpreadsheetWorkbook.tenant_id == principal["tenant_id"]).order_by(SpreadsheetWorkbook.updated_at.desc())))


@router.get("/workbooks/{workbook_id}")
def get_workbook(workbook_id: str, principal: dict = Depends(current_principal), db: Session = Depends(get_db)):
    workbook = db.scalar(select(SpreadsheetWorkbook).where(SpreadsheetWorkbook.id == workbook_id, SpreadsheetWorkbook.tenant_id == principal["tenant_id"]))
    if not workbook:
        raise HTTPException(404, "Çalışma kitabı bulunamadı.")
    return {"id": workbook.id, "name": workbook.name, "version": workbook.version,
            "workbook": json.loads(workbook.workbook_json), "protection": json.loads(workbook.protection_json)}


@router.put("/workbooks/{workbook_id}/permissions")
def set_permissions(workbook_id: str, permissions: list[CellPermission], principal: dict = Depends(require_roles("admin")), db: Session = Depends(get_db)):
    workbook = db.scalar(select(SpreadsheetWorkbook).where(SpreadsheetWorkbook.id == workbook_id, SpreadsheetWorkbook.tenant_id == principal["tenant_id"]))
    if not workbook:
        raise HTTPException(404, "Çalışma kitabı bulunamadı.")
    protection = json.loads(workbook.protection_json)
    protection["cells"] = [item.model_dump() for item in permissions]
    workbook.protection_json = json.dumps(protection, ensure_ascii=False)
    workbook.version += 1
    db.commit()
    return {"workbook_id": workbook.id, "version": workbook.version, "permissions": protection["cells"]}


@router.post("/workbooks/{workbook_id}/changes")
def change_cell(workbook_id: str, payload: SpreadsheetCellChange, principal: dict = Depends(current_principal), db: Session = Depends(get_db)):
    workbook = db.scalar(select(SpreadsheetWorkbook).where(SpreadsheetWorkbook.id == workbook_id, SpreadsheetWorkbook.tenant_id == principal["tenant_id"]))
    if not workbook:
        raise HTTPException(404, "Çalışma kitabı bulunamadı.")
    if payload.base_version != workbook.version:
        raise HTTPException(409, "Çalışma kitabı sürümü güncel değil; yeniden yükleyin.")
    if not _can_edit(json.loads(workbook.protection_json), principal["role"], payload.sheet_name, payload.cell_ref):
        raise HTTPException(403, "Bu hücre için düzenleme yetkiniz yok.")
    data = json.loads(workbook.workbook_json)
    sheet = data.setdefault(payload.sheet_name, {})
    sheet[payload.cell_ref] = payload.value
    workbook.workbook_json = json.dumps(data, ensure_ascii=False)
    workbook.version += 1
    db.add(SpreadsheetChange(tenant_id=principal["tenant_id"], workbook_id=workbook.id, user_id=principal["sub"],
                             sheet_name=payload.sheet_name, cell_ref=payload.cell_ref,
                             value_json=json.dumps(payload.value, ensure_ascii=False), version=workbook.version))
    db.commit()
    if payload.sheet_name.lower() in {"vehicles", "araçlar", "filo"} and payload.cell_ref.upper().startswith("E"):
        _arvento_request("/groups", "POST", {"name": str(payload.value), "tenant_id": principal["tenant_id"]},
                         settings=_arvento_settings(db, principal["tenant_id"]))
    return {"workbook_id": workbook.id, "version": workbook.version, "sheet_name": payload.sheet_name, "cell_ref": payload.cell_ref, "value": payload.value}


@router.post("/import")
async def import_xlsx(file: UploadFile = File(...), principal: dict = Depends(current_principal), db: Session = Depends(get_db)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(415, "Yalnızca .xlsx dosyaları kabul edilir.")
    content = await file.read()
    if len(content) > 50 * 1024 * 1024:
        raise HTTPException(413, "Dosya 50 MB sınırını aşamaz.")
    source = load_workbook(BytesIO(content), data_only=False)
    sheets = {}
    for ws in source.worksheets:
        sheets[ws.title] = {cell.coordinate: cell.value for row in ws.iter_rows() for cell in row if cell.value is not None}
    workbook = SpreadsheetWorkbook(tenant_id=principal["tenant_id"], owner_user_id=principal["sub"],
                                   name=file.filename, workbook_json=json.dumps(sheets, ensure_ascii=False))
    db.add(workbook)
    db.commit()
    db.refresh(workbook)
    return {"id": workbook.id, "name": workbook.name, "version": workbook.version, "sheets": list(sheets)}


@router.get("/workbooks/{workbook_id}/export")
def export_xlsx(workbook_id: str, principal: dict = Depends(current_principal), db: Session = Depends(get_db)):
    workbook = db.scalar(select(SpreadsheetWorkbook).where(SpreadsheetWorkbook.id == workbook_id, SpreadsheetWorkbook.tenant_id == principal["tenant_id"]))
    if not workbook:
        raise HTTPException(404, "Çalışma kitabı bulunamadı.")
    output = BytesIO()
    target = Workbook()
    target.remove(target.active)
    for sheet_name, cells in json.loads(workbook.workbook_json).items():
        ws = target.create_sheet(sheet_name)
        for ref, value in cells.items():
            ws[ref] = value
    target.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{workbook.name.rsplit(".", 1)[0]}.xlsx"'})


@router.websocket("/workbooks/{workbook_id}/sync")
async def workbook_sync(websocket: WebSocket, workbook_id: str):
    token = websocket.query_params.get("token")
    if not token:
        await websocket.close(code=4401)
        return
    try:
        principal = decode_access_token(token)
    except Exception:
        await websocket.close(code=4401)
        return
    with SessionLocal() as db:
        workbook = db.scalar(
            select(SpreadsheetWorkbook).where(
                SpreadsheetWorkbook.id == workbook_id,
                SpreadsheetWorkbook.tenant_id == principal["tenant_id"],
            )
        )
    if not workbook:
        await websocket.close(code=4404)
        return
    await websocket.accept()
    try:
        while True:
            message = await websocket.receive_json()
            await websocket.send_json({"type": "ack", "tenant_id": principal["tenant_id"], "workbook_id": workbook_id, "change": message})
    except WebSocketDisconnect:
        return

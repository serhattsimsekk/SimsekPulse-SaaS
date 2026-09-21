from __future__ import annotations

from datetime import date
from io import BytesIO
import base64

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle, Paragraph


def fleet_capacity_excel(rows: list[dict]) -> BytesIO:
    output = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Filo Kapasite"
    headers = ["Plaka", "Model", "Kapasite Ton", "Odometer KM", "Durum"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="132238")
    for row in rows:
        sheet.append([row.get("plate", ""), row.get("model", ""), row.get("capacity_ton", 0), row.get("odometer_km", 0), row.get("status", "")])
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 20
    workbook.save(output)
    output.seek(0)
    return output


def executive_pdf(summary: dict, alerts: list[dict], logo_data: str | None = None) -> BytesIO:
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=12 * mm, leftMargin=12 * mm)
    styles = getSampleStyleSheet()
    title = "⚡ ŞimşekLog Yönetici Operasyon Raporu"
    if logo_data:
        try:
            from reportlab.platypus import Image
            image = Image(BytesIO(base64.b64decode(logo_data.split(",", 1)[1])), width=28 * mm, height=28 * mm)
            story = [image, Paragraph(title, styles["Title"]), Spacer(1, 8)]
        except (ValueError, IndexError):
            story = [Paragraph(title, styles["Title"]), Spacer(1, 8)]
    else:
        story = [Paragraph(title, styles["Title"]), Spacer(1, 8)]
    story.append(Paragraph(f"Rapor tarihi: {date.today().isoformat()}", styles["Normal"]))
    story.append(Spacer(1, 12))
    story.append(Table(
        [["KPI", "Değer"], ["Araç Sayısı", str(summary.get("vehicle_count", 0))], ["Aktif Sefer", str(summary.get("active_trips", 0))], ["Taşınan Tonaj", str(summary.get("carried_ton", 0))]],
        style=TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#132238")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("PADDING", (0, 0), (-1, -1), 6)]),
    ))
    story.append(Spacer(1, 16))
    alert_rows = [["Lokasyon", "Tip", "Demoraj", "Durum"]] + [[a.get("location_name", ""), a.get("location_type", ""), str(a.get("demurrage_amount", 0)), a.get("status", "")] for a in alerts]
    story.append(Paragraph("Demoraj alarmları", styles["Heading2"]))
    story.append(Table(alert_rows or [["Kayıt yok", "", "", ""]], style=TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ed6a5a")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("PADDING", (0, 0), (-1, -1), 6)])))
    document.build(story)
    output.seek(0)
    return output
